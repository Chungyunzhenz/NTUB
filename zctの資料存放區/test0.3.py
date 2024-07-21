import cv2
import numpy as np
from PIL import Image
import pytesseract
import xml.etree.ElementTree as ET
import os
from docx import Document
from docx.shared import Inches

pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
tessdata_dir_config = '--tessdata-dir "C:\\Program Files\\Tesseract-OCR\\tessdata"'
lang = 'chi_tra+eng'  # 支持繁体中文和英文

image_path = 'G:/263_0.jpg'
image = cv2.imread(image_path)
if image is None:
    print(f"无法加载图像: {image_path}")
    exit()

# 图像预处理，增强表格线条
gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
_, binarized = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
# 识别水平和垂直线条
horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (25, 1))
vertical_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, 25))
horizontal_lines = cv2.morphologyEx(binarized, cv2.MORPH_OPEN, horizontal_kernel, iterations=2)
vertical_lines = cv2.morphologyEx(binarized, cv2.MORPH_OPEN, vertical_kernel, iterations=2)

# 合并线条
table_structure = cv2.bitwise_or(horizontal_lines, vertical_lines)
# 寻找表格边界
contours, hierarchy = cv2.findContours(table_structure, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

root = ET.Element("document")
doc = Document()

for contour in contours:
    x, y, w, h = cv2.boundingRect(contour)
    # 对于每个表格，提取ROI并进行进一步处理以识别储存格内容
    roi = gray[y:y+h, x:x+w]
    # 这里可以加入基于roi的进一步处理，例如使用pytesseract识别储存格内的文字
    # 为简化，这里仅示范如何处理和加入到WORD中
    text = pytesseract.image_to_string(roi, lang=lang, config=tessdata_dir_config)
    # 在WORD文档中添加一个表格，根据你的实际需求调整行列数
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = text

# 保存WORD文档
output_docx_path = "output_with_table.docx"
doc.save(output_docx_path)
print(f"WORD文档已保存至: {os.path.abspath(output_docx_path)}")
